"""Servicio de TABLA de Personas — filas enriquecidas con conocimiento, filtros
por hechos del Knowledge Engine, orden y exportación (CSV / Excel).

Diseñado como base reutilizable para el futuro módulo Audiencias: la generación de
filas y la exportación no dependen de la UI ni exclusivamente de Personas."""
from __future__ import annotations

import csv
import io
from datetime import datetime

from knowledge import store as knowledge_store
from people import service
from people.repository import people_repository

# Registro de columnas: key → etiqueta legible. Base compartida con Audiencias.
COLUMNS: dict[str, str] = {
    "fullName": "Nombre",
    "primaryPhone": "Teléfono",
    "primaryEmail": "Email",
    "roles": "Rol",
    "interests": "Intereses",
    "confidenceLabel": "Confianza",
    "lastContact": "Última interacción",
    "conversationsCount": "Conversaciones",
    "budget": "Presupuesto",
    "location": "Zona",
    "profession": "Profesión",
    "status": "Estado",
}
DEFAULT_COLUMNS = ["fullName", "primaryPhone", "primaryEmail", "roles", "interests", "confidenceLabel", "lastContact"]


async def build_rows(
    *, role: str | None = None, query: str | None = None,
    fact_type: str | None = None, fact_value: str | None = None,
    budget_min: float | None = None, budget_max: float | None = None,
    sort_by: str = "fullName", sort_dir: str = "asc",
) -> list[dict]:
    rows = await people_repository.list(role)
    counts = await service.owner_counts()

    if query:
        q = query.strip().lower()
        rows = [r for r in rows if q in (r.get("full_name", "") or "").lower()
                or q in (r.get("primary_phone", "") or "")
                or q in (r.get("primary_email", "") or "").lower()
                or q in (r.get("document", "") or "")]

    ids = [r["id"] for r in rows]
    facts_map = await knowledge_store.facts_for_persons(ids)

    result = []
    for r in rows:
        knowledge = knowledge_store.summarize(facts_map.get(r["id"], []))
        enrich = await service.enrich_list_item(r, counts)
        result.append({
            "id": r["id"],
            "fullName": r.get("full_name", "") or "",
            "primaryPhone": r.get("primary_phone") or r.get("phone", "") or "",
            "primaryEmail": r.get("primary_email") or r.get("email", "") or "",
            "roles": r.get("roles") or [],
            "status": r.get("status") or "active",
            "conversationsCount": enrich.get("conversationsCount", 0),
            "lastContact": enrich.get("lastContact"),
            **knowledge,
        })

    # Filtro por CUALQUIER hecho del Knowledge Engine.
    if fact_type:
        keep = set()
        fv = fact_value.strip().lower() if fact_value else None
        for pid, fs in facts_map.items():
            for f in fs:
                if f["status"] == "review" or f["type"] != fact_type:
                    continue
                if fv and str(f["value"]).strip().lower() != fv:
                    continue
                keep.add(pid)
        result = [row for row in result if row["id"] in keep]

    if budget_min is not None or budget_max is not None:
        def ok(row):
            b = row.get("budget")
            if b is None:
                return False
            try:
                bv = float(b)
            except (TypeError, ValueError):
                return False
            if budget_min is not None and bv < budget_min:
                return False
            if budget_max is not None and bv > budget_max:
                return False
            return True
        result = [row for row in result if ok(row)]

    reverse = sort_dir == "desc"

    def key(row):
        v = row.get(sort_by)
        if isinstance(v, list):
            v = ", ".join(str(x) for x in v)
        if v is None:
            return (1, "")
        return (0, v.lower() if isinstance(v, str) else v)
    try:
        result.sort(key=key, reverse=reverse)
    except TypeError:
        result.sort(key=lambda r: str(r.get(sort_by) or ""), reverse=reverse)
    return result


# --- Exportación ------------------------------------------------------------ #
def _cell(row: dict, col: str):
    v = row.get(col)
    if col == "roles" or col == "interests":
        return ", ".join(str(x) for x in (v or []))
    if col == "budget" and v is not None:
        return f"{row.get('budgetCurrency') or ''} {v}".strip()
    if col == "lastContact" and v:
        try:
            return datetime.fromisoformat(str(v).replace("Z", "+00:00"))
        except ValueError:
            return v
    return "" if v is None else v


def to_csv(rows: list[dict], columns: list[str]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([COLUMNS.get(c, c) for c in columns])
    for r in rows:
        w.writerow([
            (_cell(r, c).strftime("%Y-%m-%d %H:%M") if isinstance(_cell(r, c), datetime) else _cell(r, c))
            for c in columns
        ])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(rows: list[dict], columns: list[str]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Personas"

    header_fill = PatternFill("solid", fgColor="1E1E1E")
    header_font = Font(bold=True, color="FFFFFF")
    labels = [COLUMNS.get(c, c) for c in columns]
    ws.append(labels)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    for r in rows:
        out = []
        for c in columns:
            val = _cell(r, c)
            out.append(val)
        ws.append(out)
        if "lastContact" in columns:
            idx = columns.index("lastContact") + 1
            cell = ws.cell(row=ws.max_row, column=idx)
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"

    # Autofiltro + ancho automático.
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{ws.max_row}"
    ws.freeze_panes = "A2"
    for i, col in enumerate(columns, start=1):
        letter = get_column_letter(i)
        max_len = len(str(labels[i - 1]))
        for row in ws.iter_rows(min_col=i, max_col=i, min_row=2):
            v = row[0].value
            max_len = max(max_len, len(str(v)) if v is not None else 0)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
