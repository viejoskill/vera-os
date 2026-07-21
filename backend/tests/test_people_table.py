"""Tests for the 'Personas Table + Export' feature: /table, /facets, /export."""
import io
import os
import csv

import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://real-estate-os-10.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"


@pytest.fixture(scope="module")
def api_client():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


# ---------- /api/people/table ----------
class TestPeopleTable:
    def test_table_basic_shape(self, api_client):
        r = api_client.get(f"{API}/people/table")
        assert r.status_code == 200, r.text
        data = r.json()
        # Shape
        for key in ("rows", "total", "page", "pageSize", "columns", "defaultColumns"):
            assert key in data, f"missing key {key}"
        assert isinstance(data["rows"], list)
        assert isinstance(data["total"], int)
        assert isinstance(data["columns"], dict)
        assert isinstance(data["defaultColumns"], list)
        # Default columns must all exist in the column catalog
        for c in data["defaultColumns"]:
            assert c in data["columns"], f"default column {c} missing from catalog"
        # Row enrichment fields (may be None or '—', but keys should be present)
        if data["rows"]:
            row = data["rows"][0]
            for k in ("id", "fullName", "primaryPhone", "roles", "status",
                     "conversationsCount", "lastContact"):
                assert k in row, f"row missing enrichment key {k}"

    def test_table_pagination(self, api_client):
        r = api_client.get(f"{API}/people/table", params={"page": 1, "pageSize": 2})
        assert r.status_code == 200
        d = r.json()
        assert d["page"] == 1
        assert d["pageSize"] == 2
        assert len(d["rows"]) <= 2
        total = d["total"]
        # Page 2 (if there is data) should return different rows
        if total > 2:
            r2 = api_client.get(f"{API}/people/table", params={"page": 2, "pageSize": 2})
            ids1 = {r["id"] for r in d["rows"]}
            ids2 = {r["id"] for r in r2.json()["rows"]}
            assert ids1.isdisjoint(ids2)

    def test_table_sort_asc_desc(self, api_client):
        asc = api_client.get(f"{API}/people/table", params={"sortBy": "fullName", "sortDir": "asc"}).json()
        desc = api_client.get(f"{API}/people/table", params={"sortBy": "fullName", "sortDir": "desc"}).json()
        names_asc = [r["fullName"] for r in asc["rows"]]
        names_desc = [r["fullName"] for r in desc["rows"]]
        if len(names_asc) > 1:
            assert names_asc == sorted(names_asc, key=lambda s: s.lower())
            assert names_desc[0].lower() >= names_asc[0].lower() or names_desc == list(reversed(names_asc))

    def test_table_filter_facts_role_agrimensor(self, api_client):
        r = api_client.get(f"{API}/people/table", params=[("facts", "role:Agrimensor")])
        assert r.status_code == 200
        d = r.json()
        # Per problem statement, only Juan Pedro Andersen should match
        names = [row["fullName"] for row in d["rows"]]
        assert len(names) == 1, f"Expected exactly 1 row for role:Agrimensor, got {names}"
        assert names[0] == "Juan Pedro Andersen", f"Expected 'Juan Pedro Andersen', got {names[0]}"

    def test_table_filter_multi_facts_and(self, api_client):
        # AND: role:Agrimensor AND intent:Servicio profesional → still Juan
        r = api_client.get(f"{API}/people/table",
                           params=[("facts", "role:Agrimensor"),
                                   ("facts", "intent:Servicio profesional")])
        assert r.status_code == 200
        d = r.json()
        names = [row["fullName"] for row in d["rows"]]
        assert names == ["Juan Pedro Andersen"], f"AND filter mismatch: {names}"

    def test_table_filter_impossible_and(self, api_client):
        # AND with two roles → nobody has both → 0 results
        r = api_client.get(f"{API}/people/table",
                           params=[("facts", "role:Agrimensor"),
                                   ("facts", "role:__no_such_role__")])
        assert r.status_code == 200
        assert r.json()["total"] == 0

    def test_table_budget_min_max(self, api_client):
        # Should not error even if no rows have budget
        r = api_client.get(f"{API}/people/table", params={"budgetMin": 1, "budgetMax": 999999999})
        assert r.status_code == 200
        d = r.json()
        # every returned row must have a numeric budget in range
        for row in d["rows"]:
            b = row.get("budget")
            assert b is not None
            bv = float(b)
            assert 1 <= bv <= 999999999

    def test_table_confidence_min(self, api_client):
        r = api_client.get(f"{API}/people/table", params={"confidenceMin": 0.6})
        assert r.status_code == 200

    def test_table_contacted_within(self, api_client):
        r = api_client.get(f"{API}/people/table", params={"contactedWithin": 30})
        assert r.status_code == 200


# ---------- /api/people/facets ----------
class TestPeopleFacets:
    def test_facets_shape(self, api_client):
        r = api_client.get(f"{API}/people/facets")
        assert r.status_code == 200, r.text
        d = r.json()
        assert "filters" in d and "values" in d
        # All business filter types are declared
        for t in ("role", "relationship", "intent", "property_type", "location"):
            assert t in d["filters"], f"filter {t} missing"
            assert t in d["values"], f"values.{t} missing"
        # Given seed data, at least 'role' should include 'Agrimensor'
        assert "Agrimensor" in d["values"]["role"], f"Expected 'Agrimensor' in role facet values, got {d['values']['role']}"
        # And intent should include 'Servicio profesional'
        assert "Servicio profesional" in d["values"]["intent"], \
            f"Expected 'Servicio profesional' in intent facet values, got {d['values']['intent']}"


# ---------- /api/people/export ----------
class TestPeopleExport:
    def test_export_csv(self, api_client):
        r = api_client.get(f"{API}/people/export", params={"format": "csv"})
        assert r.status_code == 200, r.text
        # utf-8-sig BOM
        assert r.content[:3] == b"\xef\xbb\xbf", "CSV should start with UTF-8 BOM"
        text = r.content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        assert len(rows) >= 1, "CSV should have at least a header row"
        header = rows[0]
        # header labels come from column labels (human readable, Spanish)
        assert "Nombre" in header, f"CSV header missing 'Nombre': {header}"
        # Content-Disposition attachment
        cd = r.headers.get("Content-Disposition", "")
        assert "attachment" in cd and "personas_" in cd and ".csv" in cd

    def test_export_xlsx(self, api_client):
        r = api_client.get(f"{API}/people/export", params={"format": "xlsx"})
        assert r.status_code == 200, r.text
        assert r.headers.get("content-type", "").startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        # openpyxl parses without error (previously failed on tz-aware datetimes)
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        assert ws.title == "Personas"
        # Header row present
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert "Nombre" in header, f"XLSX header missing 'Nombre': {header}"
        # freeze_panes A2 and auto_filter set
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref is not None and ws.auto_filter.ref.startswith("A1:")

    def test_export_respects_columns_param(self, api_client):
        r = api_client.get(f"{API}/people/export",
                           params={"format": "csv", "columns": "fullName,primaryPhone"})
        assert r.status_code == 200
        text = r.content.decode("utf-8-sig")
        header = next(csv.reader(io.StringIO(text)))
        assert header == ["Nombre", "Teléfono"], f"Unexpected header: {header}"

    def test_export_respects_ids_selection(self, api_client):
        # Fetch table to get one id
        t = api_client.get(f"{API}/people/table", params={"pageSize": 1}).json()
        if not t["rows"]:
            pytest.skip("no rows")
        pid = t["rows"][0]["id"]
        r = api_client.get(f"{API}/people/export",
                           params={"format": "csv", "ids": pid, "columns": "fullName"})
        assert r.status_code == 200
        rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
        # header + exactly 1 selected row
        assert len(rows) == 2, f"Expected 1 data row, got: {rows}"

    def test_export_respects_facts_filter(self, api_client):
        # Only Juan Pedro Andersen should be exported for role:Agrimensor
        r = api_client.get(f"{API}/people/export",
                           params=[("format", "csv"), ("facts", "role:Agrimensor"),
                                   ("columns", "fullName")])
        assert r.status_code == 200
        rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
        assert len(rows) == 2, f"Expected header + 1 row: {rows}"
        assert rows[1][0] == "Juan Pedro Andersen"

    def test_export_xlsx_with_facts_and_columns(self, api_client):
        r = api_client.get(f"{API}/people/export",
                           params=[("format", "xlsx"), ("facts", "role:Agrimensor"),
                                   ("columns", "fullName,lastContact")])
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert header == ["Nombre", "Última interacción"]
